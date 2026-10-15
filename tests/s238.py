# s238 · Slides from text across every brand × size (checks + drawn pixels), and PowerPoint from the Export sheet
async def test(pg, ctx, logs):
    import zipfile, re, io
    from xml.dom import minidom
    st = lambda e, *a: pg.evaluate(e, *a)
    ok = 0; n = 0
    def check(name, cond, info=''):
        nonlocal ok, n
        n += 1; ok += bool(cond)
        print(('PASS ' if cond else 'FAIL ') + name + (' · ' + str(info) if info else ''))
    O = __import__('os').environ['SHOTS'] + '/'
    SAMPLE = open(__import__('os').path.join(__import__('os').environ['SITE'], 'tools', 'deck', 'samples', 'ask-ruth.md')).read()
    U = 'http://127.0.0.1:8766/?debug&template&ff=presentation'
    await pg.goto(U); await st('localStorage.clear()'); await pg.goto(U); await pg.wait_for_timeout(1500)
    await st("document.querySelector('#home-close')?.click()"); await pg.wait_for_timeout(300)
    await pg.click('.tb-content-btn[data-type="carousel"]'); await pg.wait_for_timeout(700)
    # every brand × every carousel size: generated, every check passes, and the DRAWN text sits where calculated
    for brand in ['askruth', 'moshpit', 't4r']:
        if brand != 'askruth': await pg.select_option('#brand-select', brand); await pg.wait_for_timeout(1300)
        for fmt in ['16x9', '1x1', '4x5', '3x4', '9x16']:
            await pg.select_option('#tab-format-select', f'carousel:{fmt}'); await pg.wait_for_timeout(500)
            r = await st("(t)=>{const p=__deck.plan(t);window.__p=p;return {ok:p.ok,err:p.error,n:p.slides.length,bad:p.reports.filter(x=>!x.hard).map(x=>x.issues)}}", SAMPLE)
            if not r['ok']:
                check(f'{brand} {fmt}: laid out', False, r); continue
            await st("__deck.apply(window.__p,'replace')"); await pg.wait_for_timeout(500)
            # pixels: each slide's text ink is inside the safe area and inside its calculated box
            px = await st("""(()=>{const s=__studio.state();const out=[];for(let i=0;i<s.scenes.length;i++){
                const ink=__deck.ink(i,'text');document.querySelectorAll('#scenes .clip')[i].querySelector('.frame-pick').click();
                const c=__deck.check();const t=c.rects.find(r=>r.k==='text');out.push({ink,t,sa:c.sa,ok:c.ok,iss:c.issues.map(x=>x.code)})}return out})()""")
            await pg.keyboard.press('Escape')
            inside = all(p['ink'] is None or (p['ink']['x'] >= p['sa']['x'] - 2 and p['ink']['y'] >= p['sa']['y'] - 2 and p['ink']['x'] + p['ink']['w'] <= p['sa']['x'] + p['sa']['w'] + 2 and p['ink']['y'] + p['ink']['h'] <= p['sa']['y'] + p['sa']['h'] + 2) for p in px)
            match = all(p['ink'] is None or p['t'] is None or (p['ink']['x'] >= p['t']['x'] - 4 and p['ink']['x'] + p['ink']['w'] <= p['t']['x'] + p['t']['w'] + 4 and p['ink']['y'] >= p['t']['y'] - 4 and p['ink']['y'] + p['ink']['h'] <= p['t']['y'] + p['t']['h'] + 4) for p in px)
            hard = all(not [c for c in p['iss'] if c != 'dense'] for p in px)
            check(f'{brand} {fmt}: {r["n"]} slides · checks pass · drawn text inside safe area and its box', hard and inside and match, [p['iss'] for p in px if p['iss']] or '')
        await pg.screenshot(path=O + f's238-{brand}.png')
    # PowerPoint from the Export sheet (Presentation only)
    await pg.select_option('#brand-select', 'askruth'); await pg.wait_for_timeout(1300)
    await pg.select_option('#tab-format-select', 'carousel:4x5'); await pg.wait_for_timeout(500)
    await pg.click('#btn-export'); await pg.wait_for_timeout(400)
    check('Export at 4:5: no PowerPoint option', await st("document.querySelector('#seg-filetype [data-kind=pptx]').hidden"))
    await st("document.querySelector('#sheet-close').click()"); await pg.wait_for_timeout(200)
    await pg.select_option('#tab-format-select', 'carousel:16x9'); await pg.wait_for_timeout(600)
    await st("(t)=>{const p=__deck.plan(t);__deck.apply(p,'replace')}", SAMPLE); await pg.wait_for_timeout(600)
    await pg.keyboard.press('Escape')
    await pg.click('#btn-export'); await pg.wait_for_timeout(400)
    check('Export at Presentation 16:9: PowerPoint is offered', await st("!document.querySelector('#seg-filetype [data-kind=pptx]').hidden"))
    await pg.click('#seg-filetype [data-kind="pptx"]'); await pg.wait_for_timeout(200)
    check('…with one clear button and a note on what stays editable', await st("document.querySelector('#save-go').textContent") == 'Download PowerPoint · 6 slides' and await st("!document.querySelector('#pptx-hint').hidden"))
    PNG = await st("JSON.stringify(Array.from(document.querySelectorAll('#seg-filetype button')).map(b=>[b.dataset.kind,b.hidden]))")
    async with pg.expect_download() as dl:
        await pg.click('#save-go')
    d = await dl.value
    path = O + 's238.pptx'
    await d.save_as(path)
    check('a .pptx downloads, named after the project', d.suggested_filename.endswith('-6-slides.pptx'), d.suggested_filename)
    z = zipfile.ZipFile(path)
    from pptx import Presentation
    p = Presentation(path)
    check('PowerPoint 16:9, 6 slides', abs(p.slide_width / p.slide_height - 16 / 9) < 0.01 and len(p.slides) == 6)
    texts = [[sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()] for s in p.slides]
    sc = await st("__studio.state().scenes.map(s=>[s.h1,s.body,...(s.buttons||[]).map(b=>b.text)])")
    norm = lambda t: re.sub(r'\s+', ' ', t.replace('\x0b', ' ').replace('- ', '')).strip()
    want = [norm(' '.join(x for x in s if x)) for s in sc]
    got = [norm(' '.join(t)) for t in texts]
    check('every slide’s text is real, editable PowerPoint text (same words)', got == want, [g[:40] for g, w in zip(got, want) if g != w])
    faces = sorted({r.font.name for s in p.slides for sh in s.shapes if sh.has_text_frame for para in sh.text_frame.paragraphs for r in para.runs if r.font.name})
    check('fonts are the brand’s, by weight (Bitter SemiBold, Poppins SemiBold…)', 'Bitter SemiBold' in faces and 'Poppins SemiBold' in faces, faces)
    pics = sorted({sh.name for s in p.slides for sh in s.shapes if sh.shape_type == 13})
    check('only the background and the logo are pictures here', set(pics) <= {'Background', 'Logo'}, pics)
    bad = 0
    for nm in z.namelist():
        if re.match(r'ppt/slides/slide\d+\.xml$', nm):
            dom = minidom.parseString(z.read(nm))
            for para in dom.getElementsByTagName('a:p'):
                kids = [c for c in para.childNodes if c.nodeType == 1]
                pp = [i for i, c in enumerate(kids) if c.tagName == 'a:pPr']
                if len(pp) > 1 or (pp and pp[0] != 0): bad += 1
    check('slide XML is well-formed: paragraph properties first and once', bad == 0, bad)
    check('the export sheet closed; the design is unchanged', await st("document.querySelector('#sheet').hidden") and await st("__studio.state().scenes.length") == 6)
    print(f'RESULT {ok} / {n}')
